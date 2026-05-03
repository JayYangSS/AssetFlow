from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from assetflow.domain import TEMPLATE_TRADE_TYPES
from assetflow.models import Transaction


EXPECTED_HEADERS = [
    "成交日期",
    "成交时间",
    "证券代码",
    "证券名称",
    "交易类别",
    "成交数量",
    "成交价格",
    "发生金额",
    "证券余额",
    "佣金",
    "费用",
    "重要提醒",
]


@dataclass(frozen=True)
class ExportResult:
    output_path: Path
    row_count: int
    skipped_ids: list[int]


def _to_excel_number(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def _validate_headers(ws) -> None:
    headers = [cell.value for cell in ws[1][: len(EXPECTED_HEADERS)]]
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"Template header mismatch. Expected: {EXPECTED_HEADERS}. Actual: {headers}")


def _transaction_row(tx: Transaction) -> list[object]:
    return [
        tx.trade_date.strftime("%Y%m%d"),
        tx.trade_time.strftime("%H:%M:%S") if tx.trade_time else None,
        tx.symbol,
        tx.security_name,
        TEMPLATE_TRADE_TYPES[tx.trade_type],
        _to_excel_number(tx.quantity),
        _to_excel_number(tx.price),
        _to_excel_number(tx.net_amount),
        _to_excel_number(tx.position_balance_after),
        _to_excel_number(tx.commission),
        _to_excel_number(tx.fees),
        None,
    ]


def _is_exportable(tx: Transaction) -> bool:
    return all([
        tx.trade_date,
        tx.symbol,
        tx.security_name,
        tx.trade_type in TEMPLATE_TRADE_TYPES,
        tx.quantity is not None,
        tx.price is not None,
        tx.net_amount is not None,
    ])


def export_transactions_to_template(template_path: Path, output_path: Path, transactions: list[Transaction]) -> ExportResult:
    if not template_path.exists():
        raise FileNotFoundError(template_path)
    wb = load_workbook(template_path)
    ws = wb.active
    _validate_headers(ws)

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    skipped_ids: list[int] = []
    row_count = 0
    for tx in transactions:
        if not _is_exportable(tx):
            skipped_ids.append(tx.id or -1)
            continue
        ws.append(_transaction_row(tx))
        row_count += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return ExportResult(output_path=output_path, row_count=row_count, skipped_ids=skipped_ids)
