from datetime import date, time
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from assetflow.exporters.xlsx_template import EXPECTED_HEADERS, export_transactions_to_template
from assetflow.models import Transaction


def _make_template(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(EXPECTED_HEADERS)
    ws.append(["20230512", "10:10:10（非必填项）", "000811", "示例股票", "买入", 10000, 14.82, -148244.46, "10000（非必填项）", "0.0000（非必填项）", "44.4600（非必填项）", "请删除示例行"])
    wb.save(path)
    return path


def test_export_writes_transactions_from_row_two(tmp_path: Path) -> None:
    template = _make_template(tmp_path / "导入模板.xlsx")
    output = tmp_path / "output.xlsx"
    tx = Transaction(
        broker="htsc_global",
        market="HK",
        symbol="00700",
        security_name="腾讯控股",
        trade_type="buy",
        trade_date=date(2026, 5, 1),
        trade_time=time(10, 10, 10),
        quantity=Decimal("100"),
        price=Decimal("350.1200"),
        net_amount=Decimal("-35035.0000"),
        commission=Decimal("15.0000"),
        fees=Decimal("8.0000"),
        currency="HKD",
        position_balance_after=Decimal("100"),
        source_upload_id=1,
        source_ocr_result_id=1,
        source_candidate_id=1,
        dedupe_key="xlsx",
        confidence=0.99,
    )

    result = export_transactions_to_template(template, output, [tx])

    wb = load_workbook(output)
    ws = wb["Sheet1"]
    assert result.row_count == 1
    assert [cell.value for cell in ws[1][:12]] == EXPECTED_HEADERS
    assert ws["A2"].value == "20260501"
    assert ws["C2"].value == "00700"
    assert ws["E2"].value == "买入"
    assert ws["H2"].value == -35035
    assert ws["L2"].value is None
    assert ws.max_row == 2
